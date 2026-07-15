from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
from math import ceil

import pandas as pd
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.pdfgen import canvas

from services.management_report_export import draw_bar_chart, draw_kpi_card, draw_multiline_text
from services.report_builder import get_report_builder_section_registry

PDF_MIME = "application/pdf"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _brand_value(branding: dict | None, key: str, fallback: str) -> str:
    value = (branding or {}).get(key)
    return str(value).strip() if value else fallback


def _theme_color(branding: dict | None, key: str, fallback: str) -> str:
    value = (branding or {}).get(key)
    return str(value).strip() if value else fallback


def _section_label(key: str) -> str:
    return str(get_report_builder_section_registry().get(key, {}).get("label") or key.replace("_", " ").title())


def _page_header(pdf: canvas.Canvas, title: str, summary: dict, branding: dict | None):
    filters = summary.get("filters") or {}
    primary = _theme_color(branding, "primary_color", "#1E3A8A")
    secondary = _theme_color(branding, "secondary_color", "#0F172A")
    accent = _theme_color(branding, "accent_color", "#F97316")

    pdf.setFillColor(colors.HexColor(primary))
    pdf.rect(0, 580, 792, 32, fill=True, stroke=False)
    pdf.setFillColor(colors.white)
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(20, 591, _brand_value(branding, "school_name", "EDELWEISS SCHOOL"))
    pdf.drawRightString(772, 591, f"{title} | {filters.get('academic_year_label') or 'All'}")

    pdf.setStrokeColor(colors.HexColor("#CBD5E1"))
    pdf.setLineWidth(0.5)
    pdf.line(20, 40, 772, 40)

    pdf.setFillColor(colors.HexColor(secondary))
    pdf.setFont("Helvetica", 8)
    pdf.drawString(20, 25, f"Prepared by: {_brand_value(branding, 'prepared_by', 'OperatorOS')}")
    pdf.drawRightString(772, 25, f"Page {pdf._pageNumber}")

    pdf.setFillColor(colors.HexColor(accent))


def _safe_text(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _write_lines(pdf: canvas.Canvas, lines: list[str], x: float, y: float, step: float = 12, font_size: int = 8, color: str = "#334155"):
    pdf.setFillColor(colors.HexColor(color))
    pdf.setFont("Helvetica", font_size)
    for idx, line in enumerate(lines):
        pdf.drawString(x, y - idx * step, line)


def _render_metadata_page(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, _brand_value(branding, "report_header_title", "Management Analytics Report"), summary, branding)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.setFillColor(colors.HexColor(_theme_color(branding, "secondary_color", "#0F172A")))
    pdf.drawString(50, 530, _brand_value(branding, "report_header_title", "Management Analytics Report"))
    pdf.setFont("Helvetica", 10)
    pdf.drawString(50, 510, _brand_value(branding, "report_subtitle", "Operational report builder output"))

    fields = [
        ("School", _brand_value(branding, "school_name", "-")),
        ("Foundation", _brand_value(branding, "foundation_name", "-")),
        ("Subtitle", _brand_value(branding, "report_subtitle", "-")),
        ("Template", (template or {}).get("name") or "Default"),
        ("Output", (template or {}).get("output_format") or "both"),
        ("Section Count", str(len((template or {}).get("page_order_json") or []))),
    ]
    y = 470
    for label, value in fields:
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(50, y, label)
        pdf.setFont("Helvetica", 9)
        pdf.drawString(170, y, _safe_text(value))
        y -= 22
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(50, 330, "Resolved Filters")
    _write_lines(
        pdf,
        [
            f"Academic year: {summary['filters'].get('academic_year_label') or 'All'}",
            f"Jenjang: {summary['filters'].get('jenjang_name') or 'All'}",
            f"Class: {summary['filters'].get('class_name') or 'All'}",
            f"Subject: {summary['filters'].get('subject_name') or 'All'}",
            f"Term: {summary['filters'].get('term_label') or 'All'}",
        ],
        50,
        308,
        step=14,
    )


def _render_executive_summary(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, _brand_value(branding, "report_header_title", "Management Analytics Report"), summary, branding)
    pdf.setFillColor(colors.HexColor(_theme_color(branding, "secondary_color", "#0F172A")))
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Executive Summary")
    pdf.setFont("Helvetica", 9)
    pdf.drawString(50, 512, _brand_value(branding, "report_subtitle", "Attendance, lateness, academic performance, and intervention review"))

    attendance = summary.get("attendance_summary", {})
    lateness = summary.get("lateness_by_class", [])
    grades = summary.get("grade_by_class", [])
    below = summary.get("below_kkm_alerts", [])
    interventions = summary.get("interventions_summary", {})
    hadir_pct = f"{attendance.get('status_percentages', {}).get('hadir', 0)}%"
    lateness_days = sum(int(row.get("late_days") or 0) for row in lateness)
    avg_sum = next((row.get("sumatif_average") for row in grades if row.get("sumatif_average") is not None), None)
    avg_for = next((row.get("formatif_average") for row in grades if row.get("formatif_average") is not None), None)
    open_count = interventions.get("status_counts", {}).get("open", 0)
    draw_kpi_card(pdf, 50, 430, 105, 50, "Attendance", hadir_pct, "#10B981")
    draw_kpi_card(pdf, 165, 430, 105, 50, "Lateness", str(lateness_days), "#F97316")
    draw_kpi_card(pdf, 280, 430, 105, 50, "Avg Sumatif", _safe_text(avg_sum), "#3B82F6")
    draw_kpi_card(pdf, 395, 430, 105, 50, "Avg Formatif", _safe_text(avg_for), "#A855F7")
    draw_kpi_card(pdf, 510, 430, 105, 50, "Below KKM", str(len(below)), "#EF4444")
    draw_kpi_card(pdf, 625, 430, 105, 50, "Open Interventions", str(open_count), "#EAB308")

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(50, 365, "Report Context")
    lines = [
        f"Academic year: {summary['filters'].get('academic_year_label') or 'All'}",
        f"Jenjang: {summary['filters'].get('jenjang_name') or 'All'}",
        f"Class filter: {summary['filters'].get('class_name') or 'All'}",
        f"Subject filter: {summary['filters'].get('subject_name') or 'All'}",
        f"Term filter: {summary['filters'].get('term_label') or 'All'}",
    ]
    _write_lines(pdf, lines, 50, 345, step=14)

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(390, 365, "Top Insights")
    insights = summary.get("executive_insights") or []
    if insights:
        for idx, insight in enumerate(insights[:4]):
            y = 345 - idx * 42
            pdf.setFont("Helvetica-Bold", 8)
            pdf.drawString(390, y, f"[{insight.get('category', '').upper()}] {insight.get('title', '')}")
            pdf.setFont("Helvetica", 7.5)
            draw_multiline_text(pdf, insight.get("message", ""), 390, y - 10, 320, max_lines=2, font_size=7.5, leading=9)
    else:
        pdf.setFont("Helvetica", 9)
        pdf.drawString(390, 345, "No executive insights available.")


def _render_attendance(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Attendance", summary, branding)
    attendance = summary.get("attendance_summary", {})
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Attendance")
    pdf.setFont("Helvetica", 9)
    pdf.drawString(50, 512, "Attendance status counts and percentage.")
    counts = attendance.get("status_counts", {})
    draw_kpi_card(pdf, 50, 440, 105, 50, "Hadir", str(counts.get("hadir", 0)), "#10B981")
    draw_kpi_card(pdf, 165, 440, 105, 50, "Sakit", str(counts.get("sakit", 0)), "#3B82F6")
    draw_kpi_card(pdf, 280, 440, 105, 50, "Izin", str(counts.get("izin", 0)), "#F59E0B")
    draw_kpi_card(pdf, 395, 440, 105, 50, "Alfa", str(counts.get("alfa", 0)), "#EF4444")
    draw_kpi_card(pdf, 510, 440, 105, 50, "Attendance %", f"{attendance.get('status_percentages', {}).get('hadir', 0)}%", "#14B8A6")

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(50, 385, "Attendance Summary")
    _write_lines(
        pdf,
        [
            f"Total records: {attendance.get('total_records', 0)}",
            f"Hadir: {counts.get('hadir', 0)}",
            f"Sakit: {counts.get('sakit', 0)}",
            f"Izin: {counts.get('izin', 0)}",
            f"Alfa: {counts.get('alfa', 0)}",
        ],
        50,
        365,
    )
    if template and template.get("chart_visibility_json", {}).get("attendance", False):
        labels = ["Hadir", "Sakit", "Izin", "Alfa"]
        series = [[counts.get("hadir", 0), counts.get("sakit", 0), counts.get("izin", 0), counts.get("alfa", 0)]]
        draw_bar_chart(pdf, 50, 170, 400, 160, labels, series, ["Counts"], ["#10B981"], y_max=max(10, max(series[0]) * 1.2))


def _render_lateness(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Lateness", summary, branding)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Lateness")
    rows = summary.get("lateness_by_class") or []
    total_days = sum(int(row.get("late_days") or 0) for row in rows)
    total_minutes = sum(int(row.get("late_minutes") or 0) for row in rows)
    draw_kpi_card(pdf, 50, 440, 105, 50, "Late Days", str(total_days), "#F97316")
    draw_kpi_card(pdf, 165, 440, 105, 50, "Late Minutes", str(total_minutes), "#EA580C")
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(50, 385, "Lateness by Class")
    for idx, row in enumerate(rows[:8]):
        y = 365 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('class_name')}: {row.get('late_days', 0)} days, {row.get('late_minutes', 0)} minutes")
    if template and template.get("chart_visibility_json", {}).get("lateness", False) and rows:
        draw_bar_chart(
            pdf,
            50,
            170,
            400,
            160,
            [row.get("class_name", "") for row in rows[:8]],
            [[float(row.get("late_days") or 0) for row in rows[:8]]],
            ["Late Days"],
            ["#F97316"],
            y_max=max(10, max(float(row.get("late_days") or 0) for row in rows[:8]) * 1.2),
        )


def _render_grade_class(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Grade by Class", summary, branding)
    rows = summary.get("grade_by_class") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Grade by Class")
    for idx, row in enumerate(rows[:10]):
        y = 500 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('class_name')}: Sumatif {row.get('sumatif_average')} | Formatif {row.get('formatif_average')}")
    if template and template.get("chart_visibility_json", {}).get("grade_class", False) and rows:
        draw_bar_chart(
            pdf,
            50,
            170,
            430,
            160,
            [row.get("class_name", "") for row in rows[:8]],
            [[float(row.get("sumatif_average") or 0) for row in rows[:8]], [float(row.get("formatif_average") or 0) for row in rows[:8]]],
            ["Sumatif", "Formatif"],
            ["#3B82F6", "#A855F7"],
            y_max=100,
        )


def _render_grade_subject(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Grade by Subject", summary, branding)
    rows = summary.get("grade_by_subject") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Grade by Subject")
    for idx, row in enumerate(rows[:10]):
        y = 500 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('subject_name')}: Sumatif {row.get('sumatif_average')} | Formatif {row.get('formatif_average')}")


def _render_grade_student(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Grade by Student", summary, branding)
    rows = summary.get("grade_by_student") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Grade by Student")
    for idx, row in enumerate(rows[:12]):
        y = 500 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('student_name')} / {row.get('class_name')} / {row.get('subject_name')}: {row.get('sumatif_average')} | {row.get('formatif_average')}")


def _render_below_kkm(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Below KKM", summary, branding)
    rows = summary.get("below_kkm_alerts") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Below KKM Alerts")
    for idx, row in enumerate(rows[:12]):
        y = 500 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('student_name')} / {row.get('class_name')} / {row.get('subject_name')}: {row.get('average_score')} < {row.get('kkm_threshold')}")


def _render_interventions(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Interventions", summary, branding)
    rows = (summary.get("interventions_summary") or {}).get("due_soon") or []
    counts = (summary.get("interventions_summary") or {}).get("status_counts") or {}
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Interventions")
    draw_kpi_card(pdf, 50, 440, 105, 50, "Open", str(counts.get("open", 0)), "#EF4444")
    draw_kpi_card(pdf, 165, 440, 105, 50, "In Progress", str(counts.get("in_progress", 0)), "#F97316")
    draw_kpi_card(pdf, 280, 440, 105, 50, "Monitoring", str(counts.get("monitoring", 0)), "#EAB308")
    draw_kpi_card(pdf, 395, 440, 105, 50, "Resolved", str(counts.get("resolved", 0)), "#10B981")
    for idx, row in enumerate(rows[:10]):
        y = 385 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('student_name')} / {row.get('class_name')} / {row.get('subject_name')} / {row.get('priority')}")


def _render_historical_trends(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Historical Trends", summary, branding)
    trends = summary.get("historical_trends") or {}
    trend_series = trends.get("trend_series") or {}
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Historical Trends")
    attendance_terms = (trend_series.get("attendance") or {}).get("by_term") or []
    if attendance_terms:
        pdf.setFont("Helvetica", 8)
        for idx, row in enumerate(attendance_terms[-6:]):
            pdf.drawString(50, 500 - idx * 16, f"{row.get('term_label') or row.get('period')}: {row.get('attendance_percentage')}% attendance")
    else:
        pdf.setFont("Helvetica", 9)
        pdf.drawString(50, 500, "No historical trend data available.")


def _render_forecast(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Forecast", summary, branding)
    trends = summary.get("historical_trends") or {}
    forecast_rows = trends.get("forecast_series") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Forecast")
    for idx, row in enumerate(forecast_rows[:12]):
        y = 500 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(
            50,
            y,
            f"{row.get('metric')}: {row.get('forecast_value')} ({row.get('method')}, {row.get('confidence')}, {row.get('data_sufficiency')})",
        )


def _render_intervention_impact(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Intervention Impact", summary, branding)
    impact = summary.get("intervention_impact") or {}
    impact_summary = impact.get("summary") or {}
    rows = impact.get("impact_rows") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Intervention Impact")
    draw_kpi_card(pdf, 50, 440, 105, 50, "Total", str(impact_summary.get("total_interventions", 0)), "#334155")
    draw_kpi_card(pdf, 165, 440, 105, 50, "Open", str(impact_summary.get("open_interventions", 0)), "#F97316")
    draw_kpi_card(pdf, 280, 440, 105, 50, "Overdue", str(impact_summary.get("overdue_interventions", 0)), "#EF4444")
    draw_kpi_card(pdf, 395, 440, 105, 50, "Avg Delta", _safe_text(impact_summary.get("average_score_delta")), "#3B82F6")
    draw_kpi_card(pdf, 510, 440, 105, 50, "Improved", f"{impact_summary.get('percent_improved', 0)}%", "#10B981")
    for idx, row in enumerate(rows[:10]):
        y = 385 - idx * 18
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, y, f"{row.get('student_name')} / {row.get('class_name')} / {row.get('subject_name')} / {row.get('risk_level')}")


def _render_executive_insights(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Executive Insights", summary, branding)
    insights = summary.get("executive_insights") or []
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Executive Insights")
    if not insights:
        pdf.setFont("Helvetica", 9)
        pdf.drawString(50, 500, "No executive insights available.")
        return
    for idx, insight in enumerate(insights[:10]):
        y = 500 - idx * 38
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(50, y, f"[{insight.get('severity')}] {insight.get('title')}")
        pdf.setFont("Helvetica", 8)
        draw_multiline_text(pdf, insight.get("message", ""), 50, y - 10, 680, max_lines=2, font_size=8, leading=9)


def _render_data_quality(pdf: canvas.Canvas, summary: dict, branding: dict | None, template: dict | None):
    _page_header(pdf, "Data Quality", summary, branding)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 530, "Data Quality")
    warnings = list(summary.get("warnings") or [])
    trends = summary.get("historical_trends") or {}
    warnings.extend(trends.get("warnings") or [])
    diagnostics = list(trends.get("data_quality_diagnostics") or [])
    if not warnings and not diagnostics:
        pdf.setFont("Helvetica", 9)
        pdf.drawString(50, 500, "No data quality warnings.")
        return
    for idx, warning in enumerate(warnings[:8]):
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, 500 - idx * 16, f"- {warning}")
    offset = 500 - len(warnings[:8]) * 16 - 20
    for idx, item in enumerate(diagnostics[:8]):
        pdf.drawString(50, offset - idx * 16, f"- {item.get('message')}")


SECTION_RENDERERS = {
    "executive_summary": _render_executive_summary,
    "attendance": _render_attendance,
    "lateness": _render_lateness,
    "grade_class": _render_grade_class,
    "grade_subject": _render_grade_subject,
    "grade_student": _render_grade_student,
    "below_kkm": _render_below_kkm,
    "interventions": _render_interventions,
    "historical_trends": _render_historical_trends,
    "forecast": _render_forecast,
    "intervention_impact": _render_intervention_impact,
    "executive_insights": _render_executive_insights,
    "data_quality": _render_data_quality,
    "metadata": _render_metadata_page,
}


def build_report_builder_pdf(payload: dict) -> bytes:
    summary = payload["summary_payload"]
    template = payload.get("selected_template") or {}
    branding = payload.get("branding") or {}
    sections = payload.get("resolved_sections") or ["executive_summary"]
    stream = BytesIO()
    pdf = canvas.Canvas(stream, pagesize=landscape(letter), pageCompression=0)
    for section_key in sections:
        render = SECTION_RENDERERS.get(section_key)
        if render is None:
            continue
        render(pdf, summary, branding, template)
        pdf.showPage()
    pdf.save()
    return stream.getvalue()


def _sheet_visible(template: dict | None, sheet_name: str) -> bool:
    visibility = (template or {}).get("excel_sheet_visibility_json") or {}
    return visibility.get(sheet_name, True)


def _write_sheet(writer: pd.ExcelWriter, name: str, rows: list[dict], columns: list[tuple[str, str]], *, header_row: int = 0):
    df = pd.DataFrame(rows)
    if df.empty:
        df = pd.DataFrame([{field: None for field, _ in columns}])
    df = df[[field for field, _ in columns]] if all(field in df.columns for field, _ in columns) else df.reindex(columns=[field for field, _ in columns])
    df.columns = [header for _, header in columns]
    df.to_excel(writer, sheet_name=name, index=False, startrow=header_row)
    worksheet = writer.sheets[name]
    for idx, (_, header) in enumerate(columns):
        worksheet.set_column(idx, idx, max(len(header) + 2, 14))


def _add_chart_sheet(writer: pd.ExcelWriter, payload: dict):
    template = payload.get("selected_template") or {}
    if not _sheet_visible(template, "Charts"):
        return
    workbook = writer.book
    worksheet = workbook.add_worksheet("Charts")
    worksheet.hide_gridlines(2)
    writer.sheets["Charts"] = worksheet

    chart_specs = [
        ("Attendance_Data", "Attendance %", "attendance"),
        ("Lateness_Data", "Late Days", "lateness"),
        ("Grade_Class_Data", "Sumatif Avg", "grade_class"),
        ("Below_KKM_Data", "Avg Score", "below_kkm"),
        ("Interventions_Data", "Status", "interventions"),
        ("Trend_Grades_Data", "Sumatif Avg", "historical_trends"),
        ("Intervention_Impact_Data", "Delta", "intervention_impact"),
    ]
    anchor_row = 1
    for sheet_name, value_header, section_key in chart_specs:
        if not _sheet_visible(template, sheet_name) and sheet_name not in writer.sheets:
            continue
        if not template.get("chart_visibility_json", {}).get(section_key, True):
            continue
        if sheet_name not in writer.sheets:
            continue
        data_ws = writer.sheets[sheet_name]
        chart = workbook.add_chart({"type": "column"})
        chart.add_series({
            "name": [sheet_name, 0, 1],
            "categories": [sheet_name, 1, 0, 5, 0],
            "values": [sheet_name, 1, 1, 5, 1],
        })
        chart.set_title({"name": f"{sheet_name} Chart"})
        chart.set_style(10)
        worksheet.insert_chart(anchor_row, 0, chart, {"x_scale": 1.15, "y_scale": 1.1})
        anchor_row += 15


def build_report_builder_excel(payload: dict) -> bytes:
    summary = payload["summary_payload"]
    template = payload.get("selected_template") or {}
    branding = payload.get("branding") or {}
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine="xlsxwriter") as writer:
        workbook = writer.book
        title_fmt = workbook.add_format({"bold": True, "font_size": 16, "font_name": "Arial", "font_color": _theme_color(branding, "secondary_color", "#1E293B")})
        sub_fmt = workbook.add_format({"font_size": 10, "font_name": "Arial", "font_color": "#64748B", "italic": True})
        header_fmt = workbook.add_format({"bold": True, "font_name": "Arial", "font_size": 10, "bg_color": _theme_color(branding, "primary_color", "#1E3A8A"), "font_color": "#FFFFFF", "align": "center", "valign": "vcenter"})
        cell_fmt = workbook.add_format({"font_name": "Arial", "font_size": 10, "align": "center", "valign": "vcenter"})

        def write_single_column_sheet(name: str, rows: list[str], header: str):
            if not _sheet_visible(template, name):
                return
            worksheet = workbook.add_worksheet(name)
            writer.sheets[name] = worksheet
            worksheet.write("A1", header, header_fmt)
            for idx, row in enumerate(rows, start=2):
                worksheet.write(f"A{idx}", row, cell_fmt)
            worksheet.set_column("A:A", 60)

        if _sheet_visible(template, "README"):
            readme_ws = workbook.add_worksheet("README")
            writer.sheets["README"] = readme_ws
            readme_ws.write("A1", _brand_value(branding, "report_header_title", "Management Analytics Report"), title_fmt)
            readme_ws.write("A2", _brand_value(branding, "report_subtitle", "Operational report builder output"), sub_fmt)
            readme_ws.write("A4", "Template", header_fmt)
            readme_ws.write("B4", (template or {}).get("name") or "Default", cell_fmt)
            readme_ws.write("A5", "Sections", header_fmt)
            readme_ws.write("B5", ", ".join(payload.get("resolved_sections") or []), cell_fmt)
            readme_ws.set_column("A:A", 22)
            readme_ws.set_column("B:B", 70)

        if _sheet_visible(template, "Config"):
            config_ws = workbook.add_worksheet("Config")
            writer.sheets["Config"] = config_ws
            config_ws.write("A1", "Report Branding", title_fmt)
            config_rows = [
                ("School Name", _brand_value(branding, "school_name", "")),
                ("Foundation Name", _brand_value(branding, "foundation_name", "")),
                ("Report Title", _brand_value(branding, "report_header_title", "")),
                ("Subtitle", _brand_value(branding, "report_subtitle", "")),
                ("Prepared By", _brand_value(branding, "prepared_by", "")),
                ("Primary Color", _theme_color(branding, "primary_color", "#1E3A8A")),
                ("Secondary Color", _theme_color(branding, "secondary_color", "#0F172A")),
                ("Accent Color", _theme_color(branding, "accent_color", "#F97316")),
            ]
            for idx, (label, value) in enumerate(config_rows, start=3):
                config_ws.write(f"A{idx}", label, header_fmt)
                config_ws.write(f"B{idx}", value, cell_fmt)
            config_ws.set_column("A:A", 22)
            config_ws.set_column("B:B", 40)

        if _sheet_visible(template, "Attendance_Data"):
            rows = [
                {"Period": summary["filters"].get("term_label") or summary["filters"].get("academic_year_label"), "Attendance %": summary["attendance_summary"].get("status_percentages", {}).get("hadir", 0), "Hadir": summary["attendance_summary"].get("status_counts", {}).get("hadir", 0), "Sakit": summary["attendance_summary"].get("status_counts", {}).get("sakit", 0), "Izin": summary["attendance_summary"].get("status_counts", {}).get("izin", 0), "Alfa": summary["attendance_summary"].get("status_counts", {}).get("alfa", 0)}
            ]
            _write_sheet(writer, "Attendance_Data", rows, [("Period", "Period"), ("Attendance %", "Attendance %"), ("Hadir", "Hadir"), ("Sakit", "Sakit"), ("Izin", "Izin"), ("Alfa", "Alfa")])

        if _sheet_visible(template, "Lateness_Data"):
            rows = [{"Period": row.get("class_name"), "Late Days": row.get("late_days"), "Late Minutes": row.get("late_minutes")} for row in summary.get("lateness_by_class") or []]
            _write_sheet(writer, "Lateness_Data", rows, [("Period", "Class"), ("Late Days", "Late Days"), ("Late Minutes", "Late Minutes")])

        if _sheet_visible(template, "Grade_Class_Data"):
            rows = [{"Period": row.get("class_name"), "Sumatif Avg": row.get("sumatif_average"), "Formatif Avg": row.get("formatif_average"), "Gap": None if row.get("sumatif_average") is None or row.get("formatif_average") is None else round(float(row["sumatif_average"]) - float(row["formatif_average"]), 1)} for row in summary.get("grade_by_class") or []]
            _write_sheet(writer, "Grade_Class_Data", rows, [("Period", "Class"), ("Sumatif Avg", "Sumatif Avg"), ("Formatif Avg", "Formatif Avg"), ("Gap", "Gap")])

        if _sheet_visible(template, "Grade_Subject_Data"):
            rows = [{"Period": row.get("subject_name"), "Sumatif Avg": row.get("sumatif_average"), "Formatif Avg": row.get("formatif_average"), "Gap": None if row.get("sumatif_average") is None or row.get("formatif_average") is None else round(float(row["sumatif_average"]) - float(row["formatif_average"]), 1)} for row in summary.get("grade_by_subject") or []]
            _write_sheet(writer, "Grade_Subject_Data", rows, [("Period", "Subject"), ("Sumatif Avg", "Sumatif Avg"), ("Formatif Avg", "Formatif Avg"), ("Gap", "Gap")])

        if _sheet_visible(template, "Grade_Student_Data"):
            rows = [{"Period": row.get("student_name"), "Sumatif Avg": row.get("sumatif_average"), "Formatif Avg": row.get("formatif_average"), "Gap": None if row.get("sumatif_average") is None or row.get("formatif_average") is None else round(float(row["sumatif_average"]) - float(row["formatif_average"]), 1)} for row in summary.get("grade_by_student") or []]
            _write_sheet(writer, "Grade_Student_Data", rows, [("Period", "Student"), ("Sumatif Avg", "Sumatif Avg"), ("Formatif Avg", "Formatif Avg"), ("Gap", "Gap")])

        if _sheet_visible(template, "Below_KKM_Data"):
            rows = [
                {
                    "Student": row.get("student_name"),
                    "Class": row.get("class_name"),
                    "Subject": row.get("subject_name"),
                    "Type": row.get("assessment_type"),
                    "Avg Score": row.get("average_score"),
                    "KKM Threshold": row.get("kkm_threshold"),
                    "Intervention Status": row.get("intervention_status"),
                    "Priority": row.get("intervention_priority"),
                }
                for row in summary.get("below_kkm_alerts") or []
            ]
            _write_sheet(writer, "Below_KKM_Data", rows, [("Student", "Student"), ("Class", "Class"), ("Subject", "Subject"), ("Type", "Type"), ("Avg Score", "Avg Score"), ("KKM Threshold", "KKM Threshold"), ("Intervention Status", "Intervention Status"), ("Priority", "Priority")])

        if _sheet_visible(template, "Interventions_Data"):
            interventions = (summary.get("interventions_summary") or {}).get("due_soon") or []
            rows = [{"Period": row.get("student_name"), "Status": row.get("status"), "Priority": row.get("priority"), "Follow-up": row.get("follow_up_date")} for row in interventions]
            _write_sheet(writer, "Interventions_Data", rows, [("Period", "Student"), ("Status", "Status"), ("Priority", "Priority"), ("Follow-up", "Follow-up")])

        trends = summary.get("historical_trends") or {}
        trend_series = trends.get("trend_series") or {}
        if _sheet_visible(template, "Trend_Attendance_Data"):
            rows = [{"Period": row.get("period"), "Attendance %": row.get("attendance_percentage"), "Hadir": row.get("hadir"), "Sakit": row.get("sakit"), "Izin": row.get("izin"), "Alfa": row.get("alfa")} for row in (trend_series.get("attendance") or {}).get("by_term") or []]
            _write_sheet(writer, "Trend_Attendance_Data", rows, [("Period", "Period"), ("Attendance %", "Attendance %"), ("Hadir", "Hadir"), ("Sakit", "Sakit"), ("Izin", "Izin"), ("Alfa", "Alfa")])

        if _sheet_visible(template, "Trend_Lateness_Data"):
            rows = [{"Period": row.get("period"), "Late Days": row.get("late_days"), "Late Minutes": row.get("late_minutes")} for row in (trend_series.get("lateness") or {}).get("by_term") or []]
            _write_sheet(writer, "Trend_Lateness_Data", rows, [("Period", "Period"), ("Late Days", "Late Days"), ("Late Minutes", "Late Minutes")])

        if _sheet_visible(template, "Trend_Grades_Data"):
            rows = [{"Period": row.get("period"), "Sumatif Avg": row.get("sumatif_average"), "Formatif Avg": row.get("formatif_average"), "Gap": row.get("sumatif_formatif_gap"), "Below KKM Alerts": row.get("below_kkm_alert_count"), "Threshold Source": row.get("threshold_source") or row.get("effective_threshold_source")} for row in (trend_series.get("grades") or {}).get("by_term") or []]
            _write_sheet(writer, "Trend_Grades_Data", rows, [("Period", "Period"), ("Sumatif Avg", "Sumatif Avg"), ("Formatif Avg", "Formatif Avg"), ("Gap", "Gap"), ("Below KKM Alerts", "Below KKM Alerts"), ("Threshold Source", "Threshold Source")])

        if _sheet_visible(template, "Trend_Interventions_Data"):
            rows = [{"Period": row.get("period"), "Open": row.get("open_interventions"), "Resolved": row.get("resolved_interventions"), "Overdue": row.get("overdue_followups"), "Resolution Rate": row.get("resolution_rate")} for row in (trend_series.get("interventions") or {}).get("by_term") or []]
            _write_sheet(writer, "Trend_Interventions_Data", rows, [("Period", "Period"), ("Open", "Open"), ("Resolved", "Resolved"), ("Overdue", "Overdue"), ("Resolution Rate", "Resolution Rate")])

        if _sheet_visible(template, "Forecast_Data"):
            rows = [{"Metric": row.get("metric"), "Period": row.get("period"), "Forecast Value": row.get("forecast_value"), "Method": row.get("method"), "History Points": row.get("history_points"), "Confidence": row.get("confidence"), "Data Sufficiency": row.get("data_sufficiency"), "Warning": row.get("warning")} for row in trends.get("forecast_series") or []]
            _write_sheet(writer, "Forecast_Data", rows, [("Metric", "Metric"), ("Period", "Period"), ("Forecast Value", "Forecast Value"), ("Method", "Method"), ("History Points", "History Points"), ("Confidence", "Confidence"), ("Data Sufficiency", "Data Sufficiency"), ("Warning", "Warning")])

        if _sheet_visible(template, "Trend_Insights"):
            write_single_column_sheet("Trend_Insights", [f"[{row.get('category')}] {row.get('title')}: {row.get('message')}" for row in trends.get("executive_insights") or []], "Trend Insights")

        impact = summary.get("intervention_impact") or {}
        if _sheet_visible(template, "Intervention_Impact_Data"):
            rows = [
                {
                    "ID": row.get("intervention_id"),
                    "Student": row.get("student_name"),
                    "Class": row.get("class_name"),
                    "Subject": row.get("subject_name"),
                    "Status": row.get("status"),
                    "Priority": row.get("priority"),
                    "Baseline": row.get("baseline_average"),
                    "Latest": row.get("latest_average"),
                    "Delta": row.get("score_delta"),
                    "Threshold": row.get("effective_threshold"),
                    "Moved Above KKM": row.get("moved_above_kkm"),
                    "Overdue": row.get("is_overdue"),
                    "Risk": row.get("risk_level"),
                    "Owner": row.get("owner_name"),
                }
                for row in impact.get("impact_rows") or []
            ]
            _write_sheet(
                writer,
                "Intervention_Impact_Data",
                rows,
                [("ID", "ID"), ("Student", "Student"), ("Class", "Class"), ("Subject", "Subject"), ("Status", "Status"), ("Priority", "Priority"), ("Baseline", "Baseline"), ("Latest", "Latest"), ("Delta", "Delta"), ("Threshold", "Threshold"), ("Moved Above KKM", "Moved Above KKM"), ("Overdue", "Overdue"), ("Risk", "Risk"), ("Owner", "Owner")],
            )

        if _sheet_visible(template, "Intervention_Impact_Summary"):
            rows = [{"Metric": key.replace("_", " ").title(), "Value": value} for key, value in (impact.get("summary") or {}).items() if not isinstance(value, dict)]
            _write_sheet(writer, "Intervention_Impact_Summary", rows, [("Metric", "Metric"), ("Value", "Value")])

        if _sheet_visible(template, "Risk_Students_Data"):
            rows = [{"Student": row.get("student_name"), "Class": row.get("class_name"), "Subject": row.get("subject_name"), "Risk": row.get("risk_level"), "Latest": row.get("latest_average"), "Threshold": row.get("effective_threshold"), "Overdue": row.get("is_overdue"), "Reasons": ", ".join(row.get("risk_reasons") or [])} for row in impact.get("student_risk_list") or []]
            _write_sheet(writer, "Risk_Students_Data", rows, [("Student", "Student"), ("Class", "Class"), ("Subject", "Subject"), ("Risk", "Risk"), ("Latest", "Latest"), ("Threshold", "Threshold"), ("Overdue", "Overdue"), ("Reasons", "Reasons")])

        if _sheet_visible(template, "Owner_Workload_Data"):
            rows = [{"Owner": row.get("owner_name"), "Total": row.get("total_interventions"), "Open": row.get("open_interventions"), "Resolved": row.get("resolved_interventions"), "Overdue": row.get("overdue_interventions"), "Avg Delta": row.get("average_score_delta"), "High Risk": row.get("high_risk_count")} for row in impact.get("owner_workload_summary") or []]
            _write_sheet(writer, "Owner_Workload_Data", rows, [("Owner", "Owner"), ("Total", "Total"), ("Open", "Open"), ("Resolved", "Resolved"), ("Overdue", "Overdue"), ("Avg Delta", "Avg Delta"), ("High Risk", "High Risk")])

        if _sheet_visible(template, "Insights"):
            write_single_column_sheet("Insights", [f"[{row.get('category')}] {row.get('title')}: {row.get('message')}" for row in summary.get("executive_insights") or []], "Executive Insights")

        _add_chart_sheet(writer, payload)

    return stream.getvalue()

